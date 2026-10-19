"""Excel template generation, bulk upload with row-level validation, and formatted report export."""
import io
import logging
from dataclasses import dataclass, field

import pandas as pd
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from ..forms import ForecastRowForm
from ..models import ForecastAuditLog, ManpowerForecast, ProcessMaster, whole_fte
from .forecast_service import save_forecast

logger = logging.getLogger(__name__)

UPLOAD_COLUMNS = [
    ("Function Name", "function_name"),
    ("Process Name", "process_name"),
    ("Frequency", "frequency"),
    ("Volume", "volume"),
    ("Current FTE", "current_fte"),
    ("AHT Minutes", "aht_minutes"),
    ("Working Hours Per Day", "working_hours_per_day"),
    ("Working Days Per Week", "working_days_per_week"),
    ("Working Days Per Month", "working_days_per_month"),
    ("Contingency Percentage", "contingency_percentage"),
    ("Total Time Spent Hours", "total_time_spent_hours"),
    ("Run Monte Carlo", "run_monte_carlo"),
    ("Volume Variation Percentage", "volume_variation_percentage"),
    ("AHT Variation Percentage", "time_variation_percentage"),
    ("Simulation Count", "simulation_count"),
    ("Growth Percentage", "growth_rate_percentage"),
    ("Growth Period", "growth_period"),
    ("AHT Change Percentage", "aht_change_percentage"),
    ("AHT Change Period", "aht_change_period"),
    ("Forecast Horizon Months", "forecast_horizon_months"),
    ("Remarks", "remarks"),
]
LEGACY_COLUMNS = [
    ("Average Processing Time", "avg_processing_time"),
    ("Time Unit", "time_unit"),
    ("Processing Time Variation Percentage", "time_variation_percentage"),
]
REQUIRED_HEADERS = ["Function Name", "Process Name", "Frequency", "Volume", "Current FTE"]
HEADER_BY_KEY = {key: header for header, key in UPLOAD_COLUMNS + LEGACY_COLUMNS[:2]}
HEADER_BY_KEY["time_variation_percentage"] = "AHT Variation Percentage"
FREQ_LIST = '"Daily,Weekly,Monthly,Quarterly,Half-Yearly,Yearly"'

BRAND = "1F4E79"
HEADER_FILL = PatternFill("solid", fgColor=BRAND)
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14, color=BRAND)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
WRAP = Alignment(vertical="top", wrap_text=True)
GROUP_FILLS = {"Input Values": "2E75B6", "Deterministic Results": "548235",
               "Monte Carlo Results": "7030A0", "Growth Projection": "0E7C86", "Outcome": "C55A11"}
STATUS_FILLS = {"Less Manpower": "F8D7DA", "Sufficient Manpower": "D1E7DD", "Higher Manpower": "FFE8B3"}
RAG_FILLS = {"danger": "F8D7DA", "amber": "FFE8B3", "success": "D1E7DD", "info": "CFE2FF"}
NOT_RUN = "Not run"
NO_GROWTH = "No projection"
NOT_RUN_FONT = Font(italic=True, color="808080")
CALC_FILL = PatternFill("solid", fgColor="FFF3CD")


class UploadFormatError(Exception):
    """Raised when the uploaded workbook cannot be processed at all."""


@dataclass
class UploadResult:
    total_rows: int = 0
    created_ids: list = field(default_factory=list)
    errors: list = field(default_factory=list)
    auto_filled_rows: list = field(default_factory=list)

    @property
    def success_count(self):
        return len(self.created_ids)

    @property
    def error_count(self):
        return len(self.errors)


def workbook_bytes(wb):
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def style_header_row(ws, row, first_col, last_col, fill=HEADER_FILL):
    for col in range(first_col, last_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill, cell.font, cell.alignment, cell.border = fill, HEADER_FONT, CENTER, BORDER
    ws.row_dimensions[row].height = 30


def autosize(ws, min_width=10, max_width=45):
    for col_idx in range(1, ws.max_column + 1):
        longest = 0
        for row_idx in range(1, ws.max_row + 1):
            value = ws.cell(row=row_idx, column=col_idx).value
            if value is not None:
                longest = max(longest, len(str(value)))
        ws.column_dimensions[get_column_letter(col_idx)].width = max(min_width, min(max_width, longest + 2))


def _col(header):
    return get_column_letter([h for h, _ in UPLOAD_COLUMNS].index(header) + 1)


# ------------------------------------------------------------------ Template
def generate_upload_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "Forecast Upload"
    headers = [h for h, _ in UPLOAD_COLUMNS]
    ws.append(headers)
    style_header_row(ws, 1, 1, len(headers))
    pairs = list(ProcessMaster.objects.filter(is_active=True, function__is_active=True)
                 .select_related("function").order_by("function__function_name", "process_name")[:2]) or [None]
    for i, p in enumerate(pairs):
        fname = p.function.function_name if p else "Finance"
        pname = p.process_name if p else "Invoice Processing"
        freq = p.get_default_frequency_display() if p else "Monthly"
        ws.append([fname, pname, freq, 2200, 2.5 if i == 0 else 3, 10 if i == 0 else None, 8,
                   5.5 if i == 0 else 5, 23.83 if i == 0 else 22, 15, None, "Yes" if i == 0 else "No", 10, 15, 10000,
                   5 if i == 0 else 0, "Monthly", -2 if i == 0 else 0, "Monthly", 12,
                   "Measured AHT entered" if i == 0 else "AHT blank = calculated from team working hours"])
    validations = [
        (DataValidation(type="list", formula1=FREQ_LIST, allow_blank=False, showErrorMessage=True,
                        errorTitle="Invalid frequency", error="Select a frequency from the list."), "Frequency"),
        (DataValidation(type="list", formula1='"Yes,No"', allow_blank=True, showErrorMessage=True,
                        errorTitle="Invalid value", error="Select Yes or No."), "Run Monte Carlo"),
        (DataValidation(type="list", formula1=FREQ_LIST, allow_blank=True, showErrorMessage=True,
                        errorTitle="Invalid growth period", error="Select a period from the list."), "Growth Period"),
        (DataValidation(type="list", formula1=FREQ_LIST, allow_blank=True, showErrorMessage=True,
                        errorTitle="Invalid AHT change period", error="Select a period from the list."),
         "AHT Change Period"),
        (DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", showErrorMessage=True,
                        errorTitle="Invalid FTE", error="Current FTE must be 0 or more (decimals allowed, e.g. 2.5)."),
         "Current FTE"),
    ]
    for dv, header in validations:
        ws.add_data_validation(dv)
        col = _col(header)
        dv.add(f"{col}2:{col}2000")
    ws.freeze_panes = "A2"
    autosize(ws, min_width=14)

    ins = wb.create_sheet("Instructions")
    ins["A1"] = "Bulk Forecast Upload - Instructions"
    ins["A1"].font = TITLE_FONT
    ins.append([])
    ins.append(["Column", "Required", "Rule / Default"])
    style_header_row(ins, 3, 1, 3)
    rules = [
        ("Function Name", "Yes", "Must exist and be active in Function Master"),
        ("Process Name", "Yes", "Must exist under the given function in Process Master"),
        ("Frequency", "Yes", "Daily, Weekly, Monthly, Quarterly, Half-Yearly, Yearly"),
        ("Volume", "Yes", "Greater than 0 (decimals allowed)"),
        ("Current FTE", "Yes", "0 or more, decimals allowed (e.g. 2.5)"),
        ("AHT Minutes", "No", "Measured time per item in minutes. If blank: Total Time Spent Hours is used; if that is "
                              "also blank, AHT is calculated = FTE x hours/day x working days x 60 / volume "
                              "(contingency not deducted; the status is then not a real measure)"),
        ("Working Hours Per Day", "No", "Greater than 0 and up to 24 (default 8)"),
        ("Working Days Per Week", "No", "0.5-7, decimals allowed (default 5). Mon-Fri 8 h + Sat 4 h = 5.5"),
        ("Working Days Per Month", "No", "1-31, decimals allowed (default 22). 5.5 days a week = 23.83"),
        ("Contingency Percentage", "No", "0 to below 100 (default 15). Used for Required FTE, not for calculated AHT"),
        ("Total Time Spent Hours", "No", "Alternative to AHT Minutes: hands-on hours in the period (AHT = hours x 60 / volume)"),
        ("Run Monte Carlo", "No", "Yes / No (default Yes)"),
        ("Volume Variation Percentage", "No", "0 or more (default 10). Used only when Run Monte Carlo = Yes"),
        ("AHT Variation Percentage", "No", "0 or more (default 15). Used only when Run Monte Carlo = Yes"),
        ("Simulation Count", "No", "1,000 - 100,000 (default 10,000). Used only when Run Monte Carlo = Yes"),
        ("Growth Percentage", "No", "Volume change per growth period, e.g. 5 = +5% (default 0)"),
        ("Growth Period", "No", "Daily, Weekly, Monthly, Quarterly, Half-Yearly, Yearly (default Monthly)"),
        ("AHT Change Percentage", "No", "AHT change per period, e.g. 3 = 3% slower, -5 = 5% faster (default 0)"),
        ("AHT Change Period", "No", "Daily, Weekly, Monthly, Quarterly, Half-Yearly, Yearly (default Monthly)"),
        ("Forecast Horizon Months", "No", "1-60 months (default 12)"),
        ("Remarks", "No", "Free text"),
    ]
    for r in rules:
        ins.append(list(r))
    ins.append([])
    ins.append(["Older templates with 'Average Processing Time' and 'Time Unit' are still accepted."])
    for row in ins.iter_rows(min_row=4, max_row=4 + len(rules) - 1, max_col=3):
        for cell in row:
            cell.border = BORDER
            cell.alignment = WRAP
    autosize(ins, min_width=12, max_width=90)

    ms = wb.create_sheet("Valid Masters")
    ms.append(["Function Name", "Process Name", "Default Frequency"])
    style_header_row(ms, 1, 1, 3)
    for p in ProcessMaster.objects.filter(is_active=True, function__is_active=True).select_related("function"):
        ms.append([p.function.function_name, p.process_name, p.get_default_frequency_display()])
    autosize(ms, min_width=16)
    return workbook_bytes(wb)


# ------------------------------------------------------------------ Upload
def _clean_cell(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _format_errors(form_errors):
    parts = []
    for key, messages in form_errors.items():
        label = "Row" if key == "__all__" else HEADER_BY_KEY.get(key, key)
        parts.extend(f"{label}: {m}" for m in messages)
    return "; ".join(parts)


def process_upload(uploaded_file, user) -> UploadResult:
    """Validate every row; save valid rows, collect errors for invalid rows (upload never stops midway)."""
    try:
        df = pd.read_excel(uploaded_file, sheet_name=0, dtype=object, engine="openpyxl")
    except Exception as exc:
        raise UploadFormatError(f"Unable to read the Excel file: {exc}") from exc
    df.columns = [str(c).strip() for c in df.columns]
    lower_cols = {c.lower() for c in df.columns}
    missing = [h for h in REQUIRED_HEADERS if h.lower() not in lower_cols]
    if missing:
        raise UploadFormatError("Missing required columns: " + ", ".join(missing))
    header_map = {h.lower(): key for h, key in UPLOAD_COLUMNS + LEGACY_COLUMNS}
    rename = {c: header_map[c.lower()] for c in df.columns if c.lower() in header_map}
    df = df.rename(columns=rename)[list(rename.values())].dropna(how="all")
    df = df.loc[:, ~df.columns.duplicated()]
    if df.empty:
        raise UploadFormatError("The uploaded file contains no data rows.")

    result = UploadResult(total_rows=len(df))
    seen = set()
    for idx, row in df.iterrows():
        row_number = int(idx) + 2
        data = {k: _clean_cell(v) for k, v in row.items()}
        ident = {"row_number": row_number, "function_name": data.get("function_name", ""),
                 "process_name": data.get("process_name", "")}
        form = ForecastRowForm(data)
        if not form.is_valid():
            result.errors.append({**ident, "errors": _format_errors(form.errors)})
            continue
        cd = form.cleaned_data
        key = (cd["function"].pk, cd["process"].pk)
        if key in seen:
            result.errors.append({**ident, "errors": "Duplicate Function/Process combination in this file; "
                                                     "only the first occurrence was processed."})
            continue
        seen.add(key)
        try:
            forecast = ManpowerForecast(
                function=cd["function"], process=cd["process"], frequency=cd["frequency"],
                volume=cd["volume"], current_fte=cd["current_fte"],
                total_time_spent_hours=cd["total_time_spent_hours"],
                working_hours_per_day=cd["working_hours_per_day"],
                working_days_per_week=cd["working_days_per_week"],
                working_days_per_month=cd["working_days_per_month"],
                contingency_percentage=cd["contingency_percentage"], run_monte_carlo=cd["run_monte_carlo"],
                volume_variation_percentage=cd["volume_variation_percentage"],
                time_variation_percentage=cd["time_variation_percentage"], simulation_count=cd["simulation_count"],
                growth_rate_percentage=cd["growth_rate_percentage"], growth_period=cd["growth_period"],
                aht_change_percentage=cd["aht_change_percentage"], aht_change_period=cd["aht_change_period"],
                forecast_horizon_months=cd["forecast_horizon_months"], remarks=cd.get("remarks") or "",
            )
            save_forecast(forecast, user, action=ForecastAuditLog.Action.UPLOAD)
            result.created_ids.append(forecast.pk)
            if form.total_time_auto_filled:
                result.auto_filled_rows.append(row_number)
        except Exception as exc:
            logger.exception("Bulk upload row %s failed", row_number)
            result.errors.append({**ident, "errors": f"Unexpected processing error: {exc}"})
    return result


def build_error_report(errors):
    wb = Workbook()
    ws = wb.active
    ws.title = "Upload Errors"
    ws.append(["Row Number", "Function Name", "Process Name", "Validation Errors"])
    style_header_row(ws, 1, 1, 4)
    for e in errors:
        ws.append([e["row_number"], e["function_name"], e["process_name"], e["errors"]])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=4):
        for cell in row:
            cell.border, cell.alignment = BORDER, WRAP
    autosize(ws, min_width=14, max_width=80)
    ws.freeze_panes = "A2"
    return workbook_bytes(wb)


# ------------------------------------------------------------------ Report export
def _num(value):
    return float(value) if value is not None else None


def _a(attr):
    return lambda f, s: _num(getattr(f, attr))


def _mc_input(attr):
    return lambda f, s: _num(getattr(f, attr)) if f.run_monte_carlo else NOT_RUN


def _sa(attr):
    return lambda f, s: _num(getattr(s, attr)) if s else NOT_RUN


def _g(getter):
    return lambda f, s: getter(f) if f.has_projection else NO_GROWTH


def _shortfall_text(f):
    if f.projected_shortfall_month is None:
        return "Not within horizon"
    if f.projected_shortfall_month == 0:
        return "Already short"
    return f"{f.shortfall_label} (month {f.projected_shortfall_month})"


REPORT_GROUPS = [
    ("Input Values", [
        ("Forecast ID", lambda f, s: f.pk),
        ("Function", lambda f, s: f.function.function_name),
        ("Process", lambda f, s: f.process.process_name),
        ("Frequency", lambda f, s: f.get_frequency_display()),
        ("Volume", _a("volume")),
        ("Current FTE", _a("current_fte")),
        ("AHT (min)", _a("processing_time_minutes")),
        ("AHT Source", lambda f, s: "Calculated (Calculate AHT)" if f.aht_is_calculated else "Entered (measured)"),
        ("Working Hours/Day", _a("working_hours_per_day")),
        ("Working Days/Week", _a("working_days_per_week")),
        ("Working Days/Month", _a("working_days_per_month")),
        ("Contingency %", _a("contingency_percentage")),
        ("Total Time Spent (h)", _a("total_time_spent_hours")),
        ("Monte Carlo", lambda f, s: "Yes" if f.run_monte_carlo else "No"),
        ("Volume Variation %", _mc_input("volume_variation_percentage")),
        ("AHT Variation %", _mc_input("time_variation_percentage")),
        ("Simulation Count", lambda f, s: f.simulation_count if f.run_monte_carlo else NOT_RUN),
        ("Seed", lambda f, s: f.simulation_seed if f.run_monte_carlo else None),
        ("Volume Growth %", _a("growth_rate_percentage")),
        ("Growth Period", lambda f, s: f.get_growth_period_display()),
        ("AHT Change %", _a("aht_change_percentage")),
        ("AHT Change Period", lambda f, s: f.get_aht_change_period_display()),
        ("Horizon (Months)", lambda f, s: f.forecast_horizon_months),
    ]),
    ("Deterministic Results", [
        ("Daily Volume", _a("daily_volume")),
        ("Productive Min/FTE", _a("productive_minutes_per_fte")),
        ("Workload Hours", _a("workload_hours")),
        ("Workload FTE (exact)", _a("required_fte")),
        ("Required FTE (people)", lambda f, s: f.required_people),
        ("FTE Gap", lambda f, s: float(f.people_gap)),
        ("Utilization %", _a("utilization_percentage")),
        ("Unused Capacity %", _a("unused_capacity_percentage")),
        ("Break-even AHT (min)", _a("breakeven_aht_minutes")),
        ("AHT Headroom %", _a("aht_headroom_percentage")),
    ]),
    ("Monte Carlo Results", [
        ("Average FTE", _sa("average_required_fte")),
        ("P50 FTE", _sa("p50_required_fte")),
        ("P90 FTE", _sa("p90_required_fte")),
        ("P95 FTE", _sa("p95_required_fte")),
        ("Risk-Adjusted FTE (P90, people)", lambda f, s: whole_fte(s.risk_adjusted_recommended_fte) if s else NOT_RUN),
        ("Sufficiency %", _sa("sufficiency_probability")),
        ("Failure %", _sa("failure_probability")),
    ]),
    ("Growth Projection", [
        ("Projection Drivers", _g(lambda f: f.growth_drivers_text)),
        ("Sufficient Until", _g(lambda f: f.sufficient_until_label or "-")),
        ("Shortfall From", _g(_shortfall_text)),
        ("Horizon End", _g(lambda f: f.projection_end_label)),
        ("AHT at Horizon (min)", _g(lambda f: f.projection_end_aht)),
        ("Required FTE at Horizon (people)", _g(lambda f: whole_fte(f.projected_horizon_required_fte))),
        ("FTE Needed at Horizon", _g(lambda f: whole_fte(f.projected_horizon_fte_needed))),
        ("Additional FTE at Horizon", _g(lambda f: _num(f.horizon_additional_fte))),
        ("Status at Horizon", _g(lambda f: f.get_projected_horizon_status_display())),
        ("Projection Summary", _g(lambda f: f.projection_summary)),
    ]),
    ("Outcome", [
        ("Status", lambda f, s: f.get_status_display()),
        ("Status Basis", lambda f, s: "Deterministic + Monte Carlo" if s else "Deterministic only"),
        ("High Utilization Risk", lambda f, s: "Yes" if f.high_utilization_risk else "No"),
        ("Recommendation", lambda f, s: f.recommendation),
        ("Approval Status", lambda f, s: f.get_approval_status_display()),
        ("Approved By", lambda f, s: f.approved_by.username if f.approved_by else ""),
        ("Created By", lambda f, s: f.created_by.username if f.created_by else ""),
        ("Created At", lambda f, s: timezone.localtime(f.created_at).strftime("%d-%b-%Y %H:%M")),
        ("Remarks", lambda f, s: f.remarks),
    ]),
]
PCT_LABELS = {"Contingency %", "Volume Variation %", "AHT Variation %", "Utilization %", "Unused Capacity %",
              "Sufficiency %", "Failure %", "Volume Growth %", "AHT Change %", "AHT Headroom %"}
TEXT_LABELS = {"Function", "Process", "Frequency", "Status", "Status Basis", "High Utilization Risk",
               "Recommendation", "Approval Status", "Approved By", "Created By", "Created At", "Remarks",
               "Forecast ID", "Seed", "Simulation Count", "Monte Carlo", "Required FTE (people)",
               "Risk-Adjusted FTE (P90, people)", "Required FTE at Horizon (people)", "FTE Needed at Horizon",
               "AHT Source", "Growth Period", "AHT Change Period", "Horizon (Months)", "Projection Drivers",
               "Sufficient Until", "Shortfall From", "Horizon End", "Status at Horizon", "Projection Summary"}
WRAP_LABELS = {"Recommendation", "Projection Summary"}
STATUS_LABEL_COLUMNS = ("Status", "Status at Horizon")


def _format_report_sheet(ws, labels, n_rows):
    col = 1
    for group, cols in REPORT_GROUPS:
        start, end = col, col + len(cols) - 1
        ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=end)
        cell = ws.cell(row=1, column=start, value=group)
        cell.fill = PatternFill("solid", fgColor=GROUP_FILLS[group])
        cell.font, cell.alignment = HEADER_FONT, CENTER
        col = end + 1
    style_header_row(ws, 2, 1, len(labels))
    status_cols = [labels.index(c) + 1 for c in STATUS_LABEL_COLUMNS]
    source_col = labels.index("AHT Source") + 1
    for r in range(3, n_rows + 3):
        for c, label in enumerate(labels, start=1):
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER
            if cell.value in (NOT_RUN, NO_GROWTH):
                cell.font = NOT_RUN_FONT
            elif label in WRAP_LABELS:
                cell.alignment = WRAP
            elif label in PCT_LABELS:
                cell.number_format = '0.00"%"'
            elif label not in TEXT_LABELS:
                cell.number_format = "#,##0.00"
        for sc in status_cols:
            status_cell = ws.cell(row=r, column=sc)
            fill = STATUS_FILLS.get(status_cell.value)
            if fill:
                status_cell.fill = PatternFill("solid", fgColor=fill)
                status_cell.font = Font(bold=True)
        if str(ws.cell(row=r, column=source_col).value).startswith("Calculated"):
            ws.cell(row=r, column=source_col).fill = CALC_FILL
    autosize(ws, min_width=11, max_width=45)
    ws.freeze_panes = "D3"
    if n_rows:
        ws.auto_filter.ref = f"A2:{get_column_letter(len(labels))}{n_rows + 2}"


def _write_summary_sheet(wb, forecasts):
    ws = wb.create_sheet("Summary", 0)
    ws["A1"] = "Manpower Capacity Planning - Forecast Report"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Generated on {timezone.localtime().strftime('%d-%b-%Y %H:%M')}"
    mc_count = sum(1 for f in forecasts if f.latest_summary)
    projected = [f for f in forecasts if f.has_projection]
    rows = [
        ("Total forecasts", len(forecasts)),
        ("With Monte Carlo", mc_count),
        ("Deterministic only", len(forecasts) - mc_count),
        ("AHT calculated with Calculate AHT (status not a real measure)",
         sum(1 for f in forecasts if f.aht_is_calculated)),
        ("Less Manpower (today)", sum(1 for f in forecasts if f.status == "LESS")),
        ("Sufficient Manpower (today)", sum(1 for f in forecasts if f.status == "SUFFICIENT")),
        ("Higher Manpower (today)", sum(1 for f in forecasts if f.status == "HIGHER")),
        ("High utilization risk", sum(1 for f in forecasts if f.high_utilization_risk)),
        ("Total Current FTE", float(sum(f.current_fte for f in forecasts))),
        ("Total Required FTE (people, today)", sum(f.required_people for f in forecasts)),
        ("Total Risk-Adjusted FTE*", sum(whole_fte(f.risk_adjusted_fte) for f in forecasts)),
        ("Forecasts with future projection", len(projected)),
        ("Sufficient today but short in future", sum(1 for f in projected
                                                     if f.status != "LESS" and f.projected_shortfall_month)),
        ("Additional FTE needed at horizon", float(sum(f.horizon_additional_fte or 0 for f in projected))),
    ]
    ws.append([])
    ws.append(["Metric", "Value"])
    style_header_row(ws, 4, 1, 2)
    for r in rows:
        ws.append(list(r))
    for row in ws.iter_rows(min_row=5, max_row=ws.max_row, max_col=2):
        for cell in row:
            cell.border = BORDER
    ws.append([])
    ws.append(["* ceil(P90) where Monte Carlo ran; required FTE rounded up where it did not."])
    ws.cell(row=ws.max_row, column=1).font = NOT_RUN_FONT
    ws.column_dimensions["A"].width = 62
    ws.column_dimensions["B"].width = 18


def _write_projection_sheet(wb, forecasts, title="Growth Projection"):
    ws = wb.create_sheet(title)
    headers = ["Forecast ID", "Function", "Process", "Month", "Period", "Volume Factor", "AHT Factor",
               "Workload Factor", "Volume", "Daily Volume", "AHT (min)", "Break-even AHT (min)", "Current FTE",
               "Required FTE (people)", "FTE Needed", "Additional FTE", "Utilization %", "Sufficiency %", "Status"]
    ws.append(headers)
    style_header_row(ws, 1, 1, len(headers))
    for f in forecasts:
        for p in f.projection_points:
            ws.append([f.pk, f.function.function_name, f.process.process_name, p["month"], p["label"],
                       p["growth_factor"], p.get("aht_factor", 1.0), p.get("workload_factor", p["growth_factor"]),
                       p["volume"], p["daily_volume"], p.get("aht_minutes"), p.get("breakeven_aht_minutes"),
                       float(f.current_fte), whole_fte(p["required_fte"]), whole_fte(p["fte_needed"]),
                       p["additional_fte"], p["utilization"],
                       p["sufficiency"] if p["sufficiency"] is not None else NOT_RUN, p["status_label"]])
    status_col = len(headers)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(headers)):
        for cell in row:
            cell.border = BORDER
            if cell.value == NOT_RUN:
                cell.font = NOT_RUN_FONT
            elif headers[cell.column - 1] in ("Utilization %", "Sufficiency %"):
                cell.number_format = '0.00"%"'
            elif headers[cell.column - 1] in ("Volume", "Daily Volume", "AHT (min)", "Break-even AHT (min)"):
                cell.number_format = "#,##0.00"
        status_cell = row[status_col - 1]
        if STATUS_FILLS.get(status_cell.value):
            status_cell.fill = PatternFill("solid", fgColor=STATUS_FILLS[status_cell.value])
    autosize(ws, min_width=10, max_width=30)
    ws.freeze_panes = "D2"
    return ws


def export_forecasts(queryset):
    forecasts = list(queryset.select_related("function", "process", "created_by", "approved_by")
                     .prefetch_related("simulation_summaries"))
    columns = [(label, getter) for _, cols in REPORT_GROUPS for label, getter in cols]
    labels = [label for label, _ in columns]
    records = [{label: getter(f, f.latest_summary) for label, getter in columns} for f in forecasts]
    df = pd.DataFrame(records, columns=labels)
    bio = io.BytesIO()
    with pd.ExcelWriter(bio, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Forecast Report", startrow=1, index=False)
        _format_report_sheet(writer.sheets["Forecast Report"], labels, len(df))
        _write_summary_sheet(writer.book, forecasts)
        if any(f.has_projection for f in forecasts):
            _write_projection_sheet(writer.book, [f for f in forecasts if f.has_projection])
    return bio.getvalue()


def export_forecast_detail(forecast):
    s = forecast.latest_summary
    wb = Workbook()
    ws = wb.active
    ws.title = f"Forecast {forecast.pk}"
    ws["A1"] = f"Forecast #{forecast.pk} - {forecast.function.function_name} / {forecast.process.process_name}"
    ws["A1"].font = TITLE_FONT
    row = 3
    for group, cols in REPORT_GROUPS:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        cell = ws.cell(row=row, column=1, value=group)
        cell.fill = PatternFill("solid", fgColor=GROUP_FILLS[group])
        cell.font = HEADER_FONT
        row += 1
        for label, getter in cols:
            ws.cell(row=row, column=1, value=label).font = Font(bold=True)
            value_cell = ws.cell(row=row, column=2, value=getter(forecast, s))
            value_cell.alignment = WRAP
            if value_cell.value in (NOT_RUN, NO_GROWTH):
                value_cell.font = NOT_RUN_FONT
            elif label in PCT_LABELS:
                value_cell.number_format = '0.00"%"'
            if label in STATUS_LABEL_COLUMNS and STATUS_FILLS.get(value_cell.value):
                value_cell.fill = PatternFill("solid", fgColor=STATUS_FILLS[value_cell.value])
            for c in (1, 2):
                ws.cell(row=row, column=c).border = BORDER
            row += 1
        row += 1
    ws.cell(row=row, column=1, value="Status Explanation").font = Font(bold=True)
    ws.cell(row=row, column=2, value=forecast.status_explanation).alignment = WRAP
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 80
    if forecast.has_projection:
        _write_projection_sheet(wb, [forecast], title="Month-by-Month Projection")
    return workbook_bytes(wb)


def _table(ws, start_row, headers, rows, pct_cols=(), money_cols=(), rag_col=None, rag_values=None):
    for c, h in enumerate(headers, start=1):
        ws.cell(row=start_row, column=c, value=h)
    style_header_row(ws, start_row, 1, len(headers))
    for i, values in enumerate(rows, start=1):
        for c, v in enumerate(values, start=1):
            cell = ws.cell(row=start_row + i, column=c, value=v)
            cell.border = BORDER
            if c in pct_cols:
                cell.number_format = '0.0"%"'
            elif c in money_cols:
                cell.number_format = "#,##0"
        if rag_col and rag_values:
            fill = RAG_FILLS.get(rag_values[i - 1])
            if fill:
                ws.cell(row=start_row + i, column=rag_col).fill = PatternFill("solid", fgColor=fill)
    return start_row + len(rows) + 2


def export_executive(data, filters_text=""):
    k, s, cfg, cur = data["kpis"], data["score"], data["config"], data["config"]["currency"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Executive Summary"
    ws["A1"] = "Workforce Capacity - Executive Summary"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = (f"Generated {timezone.localtime().strftime('%d-%b-%Y %H:%M')} | Outlook to {data['horizon_label']} "
                f"| Cost per FTE {cur}{cfg['cost_per_fte']:,.0f}" + (f" | {filters_text}" if filters_text else ""))
    ws["A2"].font = NOT_RUN_FONT
    kpi_rows = [
        ("Capacity Health Score (0-100)", s["score"]),
        ("  Coverage (40%)", s["coverage"]), ("  Efficiency (30%)", s["efficiency"]),
        ("  Resilience (30%)", s["resilience"]),
        ("Functions / Processes", f"{k['function_count']} / {k['process_count']}"),
        ("Current FTE", float(k["total_current"])),
        ("Required FTE today (people)", k["total_required_people"]),
        ("Organisation utilization %", float(k["org_utilization"]) if k["org_utilization"] is not None else None),
        ("Processes short today", k["short_now_count"]),
        (f"Processes short by {data['horizon_label']}", k["short_future_count"]),
        ("Processes using calculated AHT", k["calculated_aht_count"]),
        ("Hiring need today (gross FTE)", float(k["hire_now"])),
        ("Redeployable FTE", k["redeployable"]),
        ("Net hiring need today (FTE)", k["net_hire_now"]),
        (f"Net hiring need by {data['horizon_label']} (FTE)", k["net_hire_horizon"]),
        ("Annual budget for today's net hires", float(k["budget_now"])),
        (f"Annual budget for net hires by {data['horizon_label']}", float(k["budget_horizon"])),
        ("Annual cost of unused capacity", float(k["idle_cost"])),
        ("Hiring cost avoided through redeployment", float(k["redeploy_savings"])),
    ]
    row = _table(ws, 4, ["Metric", "Value"], kpi_rows)
    for r in range(5, 5 + len(kpi_rows)):
        label = str(ws.cell(row=r, column=1).value).lower()
        if "budget" in label or "cost" in label:
            ws.cell(row=r, column=2).number_format = "#,##0"
    ws.cell(row=5, column=2).fill = PatternFill("solid", fgColor=RAG_FILLS.get(s["rag"], "FFFFFF"))
    ws.cell(row=row, column=1, value="Key Insights").font = Font(bold=True, size=12, color=BRAND)
    row += 1
    for ins in data["insights"]:
        ws.cell(row=row, column=1, value=f"• {ins['text']}").alignment = WRAP
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        ws.row_dimensions[row].height = 32
        row += 1
    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 20

    sc = wb.create_sheet("Function Scorecard")
    rows = [(r["function"], r["rag_label"], r["processes"], float(r["current"]), r["required_people"],
             float(r["utilization"]) if r["utilization"] is not None else None,
             float(r["hire_now"]), r["redeployable"], whole_fte(r["needed_horizon"]),
             float(r["hire_horizon"]), r["short_now"], r["short_future"]) for r in data["scorecard"]]
    _table(sc, 1, ["Function", "RAG", "Processes", "Current FTE", "Required FTE", "Utilization %",
                   "Hire Now", "Redeployable", f"FTE Needed {data['horizon_label']}",
                   f"Hire by {data['horizon_label']}", "Short Now", "Short in Future"], rows,
           pct_cols=(6,), rag_col=2, rag_values=[r["rag"] for r in data["scorecard"]])
    autosize(sc, min_width=12)

    rk = wb.create_sheet("Top Risks")
    _table(rk, 1, ["Process", "Function", "When", "Issue", "Recommended Action"],
           [(r["process"], r["function"], r["when"], r["issue"], r["action"]) for r in data["risks"]],
           rag_col=3, rag_values=[{"orange": "amber"}.get(r["severity"], r["severity"]) for r in data["risks"]])
    autosize(rk, min_width=12, max_width=70)

    hp = wb.create_sheet("Hiring Plan")
    _table(hp, 1, ["By", "Month", "Cumulative Net Hires", "New Hires in Period", "Period Budget",
                   "Cumulative Annual Budget"],
           [(p["label"], p["month"], p["cumulative_hires"], p["new_hires"], float(p["budget"]),
             float(p["cumulative_budget"])) for p in data["hiring_plan"]], money_cols=(5, 6))
    autosize(hp, min_width=14)

    rd = wb.create_sheet("Redeployment")
    _table(rd, 1, ["Move FTE", "From Process", "From Function", "To Process", "To Function", "Type"],
           [(m["fte"], m["from_process"], m["from_function"], m["to_process"], m["to_function"],
             "Within function" if m["same_function"] else "Cross-function") for m in data["moves"]])
    autosize(rd, min_width=12)
    return workbook_bytes(wb)
