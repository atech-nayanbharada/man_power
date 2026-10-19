"""Total Time Spent: auto-filled from Working Capacity, editable, with warnings when left at capacity."""
import io
from decimal import Decimal

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from forecasting.forms import ManpowerForecastForm
from forecasting.models import ManpowerForecast, ProcessMaster
from forecasting.services import excel_service
from forecasting.services.explain import calculation_steps
from forecasting.services.forecast_service import save_forecast

from .helpers import PASSWORD, form_data, make_forecast, make_masters, make_user


class FormAutoFillTests(TestCase):
    def setUp(self):
        self.function, self.process = make_masters()

    def test_blank_total_time_is_auto_filled(self):
        form = ManpowerForecastForm(data=form_data(self.function, self.process, total_time_spent_hours=""))
        self.assertTrue(form.is_valid(), form.errors)
        self.assertTrue(form.total_time_auto_filled)
        self.assertEqual(form.cleaned_data["total_time_spent_hours"], Decimal("448.8000"))

    def test_entered_value_is_kept(self):
        form = ManpowerForecastForm(data=form_data(self.function, self.process, total_time_spent_hours="300"))
        self.assertTrue(form.is_valid(), form.errors)
        self.assertFalse(form.total_time_auto_filled)
        self.assertEqual(form.cleaned_data["total_time_spent_hours"], Decimal("300"))

    def test_zero_fte_cannot_auto_fill(self):
        form = ManpowerForecastForm(data=form_data(self.function, self.process, total_time_spent_hours="",
                                                   current_fte="0"))
        self.assertIn("cannot be auto-filled", form.errors["total_time_spent_hours"][0])

    def test_saturday_half_day_auto_fill(self):
        form = ManpowerForecastForm(data=form_data(self.function, self.process, frequency="WEEKLY", volume="400",
                                                   current_fte="4", total_time_spent_hours="",
                                                   working_days_per_week="5.5"))
        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual(form.cleaned_data["total_time_spent_hours"], Decimal("149.6000"))  # 4 x 8 x .85 x 5.5

    def test_whole_number_fte_and_decimal_days(self):
        self.assertIn("current_fte", ManpowerForecastForm(data=form_data(self.function, self.process, current_fte="2.5")).errors)
        self.assertNotIn("avg_processing_time", ManpowerForecastForm().fields)


class ServiceAutoFillTests(TestCase):
    def setUp(self):
        self.user = make_user("analyst", "Analyst")
        self.function, self.process = make_masters()

    def test_save_with_blank_total_time(self):
        f = make_forecast(self.user, self.function, self.process, total_time_spent_hours=None, run_monte_carlo=False)
        self.assertEqual(f.total_time_spent_hours, Decimal("448.8000"))
        self.assertEqual(f.avg_processing_time, Decimal("12.2400"))
        self.assertEqual(f.utilization_percentage, Decimal("100.00"))
        self.assertTrue(f.total_time_matches_capacity)

    def test_actual_hours_not_flagged(self):
        f = make_forecast(self.user, self.function, self.process, total_time_spent_hours=Decimal("300"))
        self.assertFalse(f.total_time_matches_capacity)
        self.assertEqual(f.avg_processing_time, Decimal("8.1818"))
        self.assertEqual(f.required_people, 3)

    def test_zero_fte_blank_raises(self):
        f = ManpowerForecast(function=self.function, process=self.process, frequency="MONTHLY",
                             volume=Decimal("100"), current_fte=Decimal("0"))
        with self.assertRaises(ValueError):
            save_forecast(f, self.user)

    def test_calculation_steps_show_source(self):
        auto = make_forecast(self.user, self.function, self.process, total_time_spent_hours=None, run_monte_carlo=False)
        steps = {s["label"]: s for s in calculation_steps(auto)}
        self.assertEqual(steps["Team time available in the period"]["formula"], "3 FTE x 8.00 h x (1 - 15.00%) x 22 days")
        self.assertEqual(steps["Team time available in the period"]["result"], "448.80 hours")
        self.assertEqual(steps["Total time spent (used)"]["formula"], "Auto-filled from working capacity")
        self.assertTrue(steps["Total time spent (used)"]["warn"])
        actual = make_forecast(self.user, self.function, self.process, total_time_spent_hours=Decimal("300"))
        steps = {s["label"]: s for s in calculation_steps(actual)}
        self.assertEqual(steps["Total time spent (used)"]["formula"], "Entered (actual hours, overwritten)")
        self.assertEqual(steps["AHT (calculated)"]["result"], "8.18 min per item")


class WebAutoFillTests(TestCase):
    def setUp(self):
        self.analyst = make_user("analyst", "Analyst")
        self.function, self.process = make_masters()
        self.client.login(username="analyst", password=PASSWORD)

    def test_create_blank_shows_warning(self):
        r = self.client.post(reverse("forecasting:forecast_create"),
                             form_data(self.function, self.process, total_time_spent_hours="", run_monte_carlo=None),
                             follow=True)
        self.assertContains(r, "auto-filled from working capacity")
        self.assertContains(r, "Status not reliable yet")
        self.assertContains(r, "Team time available in the period")

    def test_create_actual_no_warning(self):
        r = self.client.post(reverse("forecasting:forecast_create"),
                             form_data(self.function, self.process, total_time_spent_hours="300", run_monte_carlo=None),
                             follow=True)
        self.assertNotContains(r, "Status not reliable yet")
        self.assertContains(r, "Entered (actual hours, overwritten)")

    def test_form_page(self):
        r = self.client.get(reverse("forecasting:forecast_create"))
        self.assertContains(r, "data-total-time")
        self.assertContains(r, "Reset to auto-fill")

    def test_edit_marks_manual(self):
        f = make_forecast(self.analyst, self.function, self.process, total_time_spent_hours=Decimal("300"))
        r = self.client.get(reverse("forecasting:forecast_update", args=[f.pk]))
        self.assertContains(r, 'data-manual="true"')
        self.assertContains(r, 'value="300"')


def build_upload(rows, headers=None):
    wb = Workbook()
    ws = wb.active
    ws.append(headers or [h for h, _ in excel_service.UPLOAD_COLUMNS])
    for r in rows:
        ws.append(r)
    bio = io.BytesIO()
    wb.save(bio)
    return SimpleUploadedFile("upload.xlsx", bio.getvalue())


def row(process="Invoice Processing", fte=3, hours=300, wk=5, function="Finance"):
    return [function, process, "Monthly", 2200, fte, 8, wk, 22, 15, hours, "No", None, None, None,
            0, "Monthly", 0, "Monthly", 12, ""]


class UploadAutoFillTests(TestCase):
    def setUp(self):
        self.user = make_user("analyst", "Analyst")
        self.function, self.process = make_masters()
        ProcessMaster.objects.create(function=self.function, process_name="Vendor Payments")
        ProcessMaster.objects.create(function=self.function, process_name="GST")

    def test_blank_rows_auto_filled_and_listed(self):
        result = excel_service.process_upload(build_upload([
            row(), row(process="Vendor Payments", hours=None), row(process="GST", hours=None, fte=0),
        ]), self.user)
        self.assertEqual(result.success_count, 2)
        self.assertEqual(result.auto_filled_rows, [3])
        self.assertIn("cannot be auto-filled", result.errors[0]["errors"])
        f = ManpowerForecast.objects.get(process__process_name="Vendor Payments")
        self.assertEqual(f.total_time_spent_hours, Decimal("448.8000"))

    def test_legacy_aht_column(self):
        headers = ["Function Name", "Process Name", "Frequency", "Volume", "Current FTE", "Average Processing Time", "Time Unit"]
        result = excel_service.process_upload(build_upload([["Finance", "Invoice Processing", "Monthly", 2200, 3, 10, "Minutes"]],
                                                           headers=headers), self.user)
        self.assertEqual(result.success_count, 1, result.errors)
        self.assertEqual(ManpowerForecast.objects.get().total_time_spent_hours, Decimal("366.6667"))

    def test_report_shows_source(self):
        make_forecast(self.user, self.function, self.process, total_time_spent_hours=None)
        report = load_workbook(io.BytesIO(excel_service.export_forecasts(ManpowerForecast.objects.all())))
        ws = report["Forecast Report"]
        labels = [c.value for c in ws[2]]
        self.assertEqual(ws.cell(row=3, column=labels.index("Total Time Source") + 1).value, "Auto (working capacity)")
        self.assertEqual(ws.cell(row=3, column=labels.index("Team Time Available (h)") + 1).value, 448.8)
        tpl = load_workbook(io.BytesIO(excel_service.generate_upload_template()))
        self.assertIn("Total Time Spent Hours", [c.value for c in tpl["Forecast Upload"][1]])
