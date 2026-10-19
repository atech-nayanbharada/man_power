"""AHT field + Calculate AHT button (server side: blank AHT is calculated with the same formula)."""
import io
from decimal import Decimal

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from forecasting.forms import ManpowerForecastForm
from forecasting.models import ManpowerForecast, ProcessMaster
from forecasting.services import excel_service
from forecasting.services.explain import calculation_steps
from forecasting.services.forecast_service import save_forecast

from .helpers import PASSWORD, form_data, make_forecast, make_masters, make_user


class AhtFormTests(TestCase):
    def setUp(self):
        self.function, self.process = make_masters()

    def test_fields(self):
        fields = ManpowerForecastForm().fields
        self.assertIn("avg_processing_time", fields)
        self.assertNotIn("total_time_spent_hours", fields)

    def test_entered_aht_is_kept(self):
        form = ManpowerForecastForm(data=form_data(self.function, self.process))
        self.assertTrue(form.is_valid(), form.errors)
        self.assertFalse(form.aht_auto_filled)
        self.assertEqual(form.cleaned_data["avg_processing_time"], Decimal("10"))

    def test_blank_aht_is_calculated(self):
        form = ManpowerForecastForm(data=form_data(self.function, self.process, avg_processing_time=""))
        self.assertTrue(form.is_valid(), form.errors)
        self.assertTrue(form.aht_auto_filled)
        self.assertEqual(form.cleaned_data["avg_processing_time"], Decimal("14.4000"))  # 3 x 8 x 22 = 528 h x 60 / 2200

    def test_user_example_yearly(self):
        form = ManpowerForecastForm(data=form_data(self.function, self.process, frequency="YEARLY", volume="6305",
                                                   current_fte="180", avg_processing_time=""))
        self.assertTrue(form.is_valid(), form.errors)
        # 180 x 8 x 264 = 380,160 h; x 60 / 6,305 (contingency not deducted)
        self.assertEqual(form.cleaned_data["avg_processing_time"], Decimal("3617.7002"))

    def test_contingency_not_used_for_aht(self):
        a = ManpowerForecastForm(data=form_data(self.function, self.process, avg_processing_time="",
                                                contingency_percentage="0"))
        b = ManpowerForecastForm(data=form_data(self.function, self.process, avg_processing_time="",
                                                contingency_percentage="30"))
        self.assertTrue(a.is_valid() and b.is_valid())
        self.assertEqual(a.cleaned_data["avg_processing_time"], b.cleaned_data["avg_processing_time"])

    def test_zero_fte_needs_entered_aht(self):
        form = ManpowerForecastForm(data=form_data(self.function, self.process, avg_processing_time="", current_fte="0"))
        self.assertIn("cannot be calculated", form.errors["avg_processing_time"][0])
        form = ManpowerForecastForm(data=form_data(self.function, self.process, current_fte="0"))
        self.assertTrue(form.is_valid(), form.errors)

    def test_negative_aht(self):
        form = ManpowerForecastForm(data=form_data(self.function, self.process, avg_processing_time="-1"))
        self.assertIn("avg_processing_time", form.errors)


class AhtServiceTests(TestCase):
    def setUp(self):
        self.user = make_user("analyst", "Analyst")
        self.function, self.process = make_masters()

    def test_entered_aht(self):
        f = ManpowerForecast(function=self.function, process=self.process, frequency="MONTHLY", volume=Decimal("2200"),
                             current_fte=Decimal("3"), avg_processing_time=Decimal("8"), run_monte_carlo=False)
        save_forecast(f, self.user)
        self.assertEqual(f.avg_processing_time, Decimal("8.0000"))
        self.assertEqual(f.total_time_spent_hours, Decimal("293.3333"))
        self.assertFalse(f.total_time_matches_capacity)

    def test_capacity_aht(self):
        f = make_forecast(self.user, self.function, self.process, total_time_spent_hours=None, run_monte_carlo=False)
        self.assertEqual(f.avg_processing_time, Decimal("14.4000"))
        self.assertEqual(f.utilization_percentage, Decimal("117.65"))
        self.assertEqual(f.status, "LESS")
        self.assertTrue(f.total_time_matches_capacity)

    def test_zero_fte_without_aht_raises(self):
        f = ManpowerForecast(function=self.function, process=self.process, frequency="MONTHLY",
                             volume=Decimal("100"), current_fte=Decimal("0"))
        with self.assertRaises(ValueError):
            save_forecast(f, self.user)

    def test_calculation_steps(self):
        auto = make_forecast(self.user, self.function, self.process, total_time_spent_hours=None, run_monte_carlo=False)
        steps = {s["label"]: s for s in calculation_steps(auto)}
        self.assertEqual(steps["Team time available in the period"]["formula"], "3 FTE x 8.00 h x 22 days (contingency not deducted)")
        self.assertEqual(steps["AHT (calculated with Calculate AHT)"]["formula"], "528.00 h x 60 / 2,200.00 items")
        self.assertEqual(steps["AHT (calculated with Calculate AHT)"]["result"], "14.40 min per item")
        entered = make_forecast(self.user, self.function, self.process, total_time_spent_hours=Decimal("300"))
        steps = {s["label"]: s for s in calculation_steps(entered)}
        self.assertEqual(steps["AHT (entered)"]["result"], "8.18 min per item")
        self.assertEqual(steps["Implied total time spent"]["result"], "300.00 hours")


class AhtWebTests(TestCase):
    def setUp(self):
        self.analyst = make_user("analyst", "Analyst")
        self.function, self.process = make_masters()
        self.client.login(username="analyst", password=PASSWORD)

    def test_form_page_has_button(self):
        r = self.client.get(reverse("forecasting:forecast_create"))
        self.assertContains(r, "data-aht-calc")
        self.assertContains(r, "Calculate AHT")

    def test_create_with_calculated_aht(self):
        r = self.client.post(reverse("forecasting:forecast_create"),
                             form_data(self.function, self.process, avg_processing_time="", run_monte_carlo=None),
                             follow=True)
        self.assertContains(r, "Status not reliable yet")
        self.assertContains(r, "AHT (calculated with Calculate AHT)")

    def test_create_with_entered_aht(self):
        r = self.client.post(reverse("forecasting:forecast_create"),
                             form_data(self.function, self.process, avg_processing_time="8", run_monte_carlo=None),
                             follow=True)
        self.assertNotContains(r, "Status not reliable yet")
        self.assertContains(r, "AHT (entered)")

    def test_edit_page(self):
        auto = make_forecast(self.analyst, self.function, self.process, total_time_spent_hours=None, run_monte_carlo=False)
        self.assertContains(self.client.get(reverse("forecasting:forecast_update", args=[auto.pk])), 'data-calculated="true"')
        entered = make_forecast(self.analyst, self.function, self.process, total_time_spent_hours=Decimal("300"))
        self.assertContains(self.client.get(reverse("forecasting:forecast_update", args=[entered.pk])), 'value="8.1818"')


def build_upload(rows, headers=None):
    wb = Workbook()
    ws = wb.active
    ws.append(headers or [h for h, _ in excel_service.UPLOAD_COLUMNS])
    for r in rows:
        ws.append(r)
    bio = io.BytesIO()
    wb.save(bio)
    return SimpleUploadedFile("upload.xlsx", bio.getvalue())


def row(process="Invoice Processing", fte=3, hours=300):
    return ["Finance", process, "Monthly", 2200, fte, 8, 5, 22, 15, hours, "No", None, None, None,
            0, "Monthly", 0, "Monthly", 12, ""]


class UploadTests(TestCase):
    def setUp(self):
        self.user = make_user("analyst", "Analyst")
        self.function, self.process = make_masters()
        ProcessMaster.objects.create(function=self.function, process_name="Vendor Payments")
        ProcessMaster.objects.create(function=self.function, process_name="GST")

    def test_upload(self):
        result = excel_service.process_upload(build_upload([
            row(), row(process="Vendor Payments", hours=None), row(process="GST", hours=None, fte=0),
        ]), self.user)
        self.assertEqual(result.success_count, 2)
        self.assertEqual(result.auto_filled_rows, [3])
        self.assertEqual(ManpowerForecast.objects.get(process=self.process).avg_processing_time, Decimal("8.1818"))
        self.assertEqual(ManpowerForecast.objects.get(process__process_name="Vendor Payments").avg_processing_time,
                         Decimal("14.4000"))

    def test_legacy_aht_column(self):
        headers = ["Function Name", "Process Name", "Frequency", "Volume", "Current FTE", "Average Processing Time", "Time Unit"]
        result = excel_service.process_upload(build_upload([["Finance", "Invoice Processing", "Monthly", 2200, 3, 10, "Minutes"]],
                                                           headers=headers), self.user)
        self.assertEqual(result.success_count, 1, result.errors)
        self.assertEqual(ManpowerForecast.objects.get().avg_processing_time, Decimal("10.0000"))

    def test_report(self):
        make_forecast(self.user, self.function, self.process, total_time_spent_hours=None)
        ws = load_workbook(io.BytesIO(excel_service.export_forecasts(ManpowerForecast.objects.all())))["Forecast Report"]
        labels = [c.value for c in ws[2]]
        self.assertEqual(ws.cell(row=3, column=labels.index("AHT (min, calculated)") + 1).value, 14.4)
