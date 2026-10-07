"""Executive (CEO) analytics: KPIs, health score, redeployment, hiring plan, risks, page and export."""
import io
from datetime import date
from decimal import Decimal

from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from forecasting.models import ApprovalStatus, FunctionMaster, ManpowerForecast, ProcessMaster
from forecasting.services.executive_service import build_executive, money, rag_for_score

from .helpers import PASSWORD, make_forecast, make_user

START = date(2026, 10, 5)


class ExecutiveDataMixin:
    """
    Four deterministic processes (Monte Carlo off so every number is exact):
      Finance / Bank Rec      150/day x 6 min, 2 FTE  -> 2.21 req, needs 3  -> LESS, hire 1
      Finance / Vendor Pay    400/wk x 12 min, 4 FTE  -> 2.35 req, needs 3  -> HIGHER (58.8%), 1 redeployable
      Finance / Invoices      2200/mo x 10 min, 3 FTE -> 2.45 req, +5%/month -> short from month 5, needs 5 at m12
      HR / Leave              300/day x 6 min, 2 FTE  -> 4.41 req, needs 5  -> LESS, hire 3
    """

    def build_data(self):
        self.user = make_user("analyst", "Analyst")
        self.finance = FunctionMaster.objects.create(function_name="Finance")
        self.hr = FunctionMaster.objects.create(function_name="HR")
        mk = lambda fn, name: ProcessMaster.objects.create(function=fn, process_name=name)  # noqa: E731
        common = dict(run_monte_carlo=False, simulation_seed=None)
        self.bank = make_forecast(self.user, self.finance, mk(self.finance, "Bank Rec"), frequency="DAILY",
                                  volume=Decimal("150"), current_fte=Decimal("2"), avg_processing_time=Decimal("6"),
                                  **common)
        self.vendor = make_forecast(self.user, self.finance, mk(self.finance, "Vendor Pay"), frequency="WEEKLY",
                                    volume=Decimal("400"), current_fte=Decimal("4"),
                                    avg_processing_time=Decimal("12"), **common)
        self.invoice = make_forecast(self.user, self.finance, mk(self.finance, "Invoices"),
                                     growth_rate_percentage=Decimal("5"), **common)
        self.leave = make_forecast(self.user, self.hr, mk(self.hr, "Leave"), frequency="DAILY",
                                   volume=Decimal("300"), current_fte=Decimal("2"),
                                   avg_processing_time=Decimal("6"), **common)

    def run_exec(self, qs=None, **kw):
        return build_executive(qs if qs is not None else ManpowerForecast.objects.all(), START,
                               cost_per_fte=kw.pop("cost", 600000), horizon=kw.pop("horizon", 12))


class ExecutiveServiceTests(ExecutiveDataMixin, TestCase):
    def setUp(self):
        self.build_data()
        self.data = self.run_exec()

    def test_statuses_of_fixture(self):
        self.assertEqual((self.bank.status, self.vendor.status, self.invoice.status, self.leave.status),
                         ("LESS", "HIGHER", "SUFFICIENT", "LESS"))

    def test_headline_kpis(self):
        k = self.data["kpis"]
        self.assertEqual(k["process_count"], 4)
        self.assertEqual(k["function_count"], 2)
        self.assertEqual(k["total_current"], Decimal("11.00"))
        self.assertEqual(k["total_required"], Decimal("11.42"))
        self.assertEqual(k["org_utilization"], Decimal("103.82"))
        self.assertEqual(k["hire_now"], Decimal("4"))
        self.assertEqual(k["redeployable"], 1)
        self.assertEqual(k["net_hire_now"], 3)
        self.assertEqual(k["hire_horizon"], Decimal("6"))
        self.assertEqual(k["net_hire_horizon"], 5)
        self.assertEqual(k["short_now_count"], 2)
        self.assertEqual(k["short_future_count"], 1)
        self.assertEqual(k["total_needed_horizon"], Decimal("16"))
        self.assertEqual(k["aht_change_count"], 0)

    def test_costs(self):
        k = self.data["kpis"]
        self.assertEqual(k["budget_now"], Decimal("1800000"))
        self.assertEqual(k["budget_horizon"], Decimal("3000000"))
        self.assertEqual(k["idle_fte"], Decimal("1.00"))
        self.assertEqual(k["idle_cost"], Decimal("600000.00"))
        self.assertEqual(k["redeploy_savings"], Decimal("600000"))
        self.assertEqual(self.data["money"]["budget_now"], "₹18.0 L")

    def test_health_score(self):
        s = self.data["score"]
        self.assertEqual(s["coverage"], 50.0)
        self.assertEqual(s["efficiency"], 52.4)
        self.assertEqual(s["resilience"], 50.0)
        self.assertEqual(s["score"], 51)
        self.assertEqual(s["rag"], "amber")

    def test_redeployment_prefers_same_function(self):
        self.assertEqual(self.data["moves"], [{
            "from_process": "Vendor Pay", "from_function": "Finance", "to_process": "Bank Rec",
            "to_function": "Finance", "fte": 1, "same_function": True}])

    def test_quarterly_hiring_plan(self):
        plan = [(p["label"], p["cumulative_hires"], p["new_hires"]) for p in self.data["hiring_plan"]]
        self.assertEqual(plan, [("Now", 3, 3), ("Jan 2027", 3, 0), ("Apr 2027", 4, 1),
                                ("Jul 2027", 4, 0), ("Oct 2027", 5, 1)])

    def test_top_risks_order(self):
        risks = self.data["risks"]
        self.assertEqual([r["process"] for r in risks], ["Leave", "Bank Rec", "Invoices"])
        self.assertEqual(risks[2]["severity"], "orange")
        self.assertEqual(risks[2]["when"], "Mar 2027")
        self.assertIn("+5% volume per month", risks[2]["issue"])

    def test_function_scorecard(self):
        rows = {r["function"]: r for r in self.data["scorecard"]}
        self.assertEqual(rows["Finance"]["current"], Decimal("9.00"))
        self.assertEqual(rows["Finance"]["hire_now"], Decimal("1"))
        self.assertEqual(rows["Finance"]["redeployable"], 1)
        self.assertEqual(rows["Finance"]["rag"], "danger")
        self.assertEqual(rows["HR"]["hire_now"], Decimal("3"))

    def test_insights(self):
        texts = " ".join(i["text"] for i in self.data["insights"])
        self.assertIn("2 of 4 processes are short of manpower today", texts)
        self.assertIn("1 FTE sit in over-staffed processes", texts)
        self.assertIn("Net hiring need by Oct 2027: 5 FTE", texts)
        self.assertIn("above the 90% safe limit", texts)
        self.assertIn("HR needs the most people today", texts)

    def test_trajectory_series(self):
        t = self.data["chart_data"]["trajectory"]
        self.assertEqual(len(t["labels"]), 13)
        self.assertEqual(t["needed"][0], 14.0)
        self.assertEqual(t["needed"][-1], 16.0)

    def test_aht_change_drives_outlook(self):
        # Vendor Pay gets +10% AHT per quarter -> 2.35 x 1.1^4 = 3.44 FTE -> needs 4 at month 12
        make_forecast(self.user, self.finance, self.vendor.process, frequency="WEEKLY", volume=Decimal("400"),
                      current_fte=Decimal("4"), avg_processing_time=Decimal("12"), run_monte_carlo=False,
                      aht_change_percentage=Decimal("10"), aht_change_period="QUARTERLY")
        d = self.run_exec()
        self.assertEqual(d["kpis"]["aht_change_count"], 1)
        self.assertEqual(d["kpis"]["total_needed_horizon"], Decimal("17"))
        self.assertTrue(any("AHT change" in i["text"] for i in d["insights"]))

    def test_shorter_horizon_and_cost(self):
        d = self.run_exec(horizon=3, cost=1000000)
        self.assertEqual(d["kpis"]["short_future_count"], 0)
        self.assertEqual(d["kpis"]["budget_now"], Decimal("3000000"))
        self.assertEqual(len(d["hiring_plan"]), 2)

    def test_function_filter_and_empty(self):
        d = self.run_exec(ManpowerForecast.objects.filter(function=self.hr))
        self.assertEqual(d["kpis"]["process_count"], 1)
        empty = self.run_exec(ManpowerForecast.objects.none())
        self.assertEqual(empty["kpis"]["process_count"], 0)
        self.assertEqual(empty["insights"], [])

    def test_helpers(self):
        self.assertEqual(rag_for_score(80), "success")
        self.assertEqual(rag_for_score(60), "amber")
        self.assertEqual(rag_for_score(40), "danger")
        self.assertEqual(money(Decimal("25000000"), "₹"), "₹2.50 Cr")
        self.assertEqual(money(Decimal("450000"), "₹"), "₹4.5 L")
        self.assertEqual(money(Decimal("1500000"), "$"), "$1.50M")


class ExecutiveViewTests(ExecutiveDataMixin, TestCase):
    def setUp(self):
        self.build_data()
        self.viewer = make_user("viewer", "Viewer")
        self.admin = make_user("boss", "Admin")

    def test_page_renders_for_all_roles(self):
        for user in (self.user, self.viewer, self.admin):
            self.client.login(username=user.username, password=PASSWORD)
            self.assertEqual(self.client.get(reverse("forecasting:executive")).status_code, 200)
            self.client.logout()

    def test_page_content(self):
        self.client.login(username=self.admin.username, password=PASSWORD)
        r = self.client.get(reverse("forecasting:executive"))
        for text in ("Capacity Health Score", "The Bottom Line", "Function Scorecard", "Top Risks",
                     "Quarterly Hiring Plan", "Redeployment Opportunities", "Key Insights", '"trajectory"'):
            self.assertContains(r, text)

    def test_viewer_sees_approved_only(self):
        self.client.login(username=self.viewer.username, password=PASSWORD)
        self.assertContains(self.client.get(reverse("forecasting:executive")), "No forecasts available")
        ManpowerForecast.objects.filter(pk=self.leave.pk).update(approval_status=ApprovalStatus.APPROVED)
        self.assertEqual(self.client.get(reverse("forecasting:executive")).context["kpis"]["process_count"], 1)

    def test_invalid_inputs_fall_back_to_defaults(self):
        self.client.login(username=self.admin.username, password=PASSWORD)
        r = self.client.get(reverse("forecasting:executive") + "?horizon=abc&cost_per_fte=-5")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.context["config"]["horizon"], 12)

    def test_excel_export(self):
        self.client.login(username=self.admin.username, password=PASSWORD)
        r = self.client.get(reverse("forecasting:executive_export"))
        wb = load_workbook(io.BytesIO(r.content))
        self.assertEqual(wb.sheetnames, ["Executive Summary", "Function Scorecard", "Top Risks", "Hiring Plan",
                                         "Redeployment"])
        self.assertEqual(wb["Executive Summary"]["B5"].value, 51)
