import io
from decimal import Decimal

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from openpyxl import Workbook, load_workbook

from forecasting.forms import ManpowerForecastForm
from forecasting.models import FunctionMaster, ManpowerForecast, ProcessMaster
from forecasting.services import excel_service
from forecasting.services.explain import calculation_steps
from forecasting.services.forecast_service import build_input, save_forecast
from forecasting.services.scenario_service import compare_scenarios

from .helpers import form_data, make_forecast, make_masters, make_user


class ForecastFormValidationTests(TestCase):
    def setUp(self):
        self.function, self.process = make_masters()

    def data(self, **kw):
        return form_data(self.function, self.process, **kw)

    def test_valid_and_no_aht_field(self):
        form = ManpowerForecastForm(data=self.data())
        self.assertTrue(form.is_valid(), form.errors)
        self.assertNotIn("avg_processing_time", form.fields)
        self.assertNotIn("time_unit", form.fields)
        self.assertIn("total_time_spent_hours", form.fields)

    def test_total_time_required(self):
        form = ManpowerForecastForm(data=self.data(total_time_spent_hours=""))
        self.assertIn("total_time_spent_hours", form.errors)

    def test_whole_number_fte(self):
        form = ManpowerForecastForm(data=self.data(current_fte="2.5"))
        self.assertIn("whole number", form.errors["current_fte"][0])
        self.assertTrue(ManpowerForecastForm(data=self.data(current_fte="0")).is_valid())

    def test_decimal_working_days(self):
        form = ManpowerForecastForm(data=self.data(working_days_per_week="5.5", working_days_per_month="23.83"))
        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual(form.cleaned_data["working_days_per_week"], Decimal("5.5"))

    def test_out_of_range_values(self):
        cases = {"volume": "0", "current_fte": "-1", "total_time_spent_hours": "0", "working_hours_per_day": "25",
                 "working_days_per_week": "7.5", "working_days_per_month": "32", "contingency_percentage": "100",
                 "simulation_count": "500", "growth_rate_percentage": "-100", "aht_change_percentage": "-100",
                 "forecast_horizon_months": "61"}
        for field, value in cases.items():
            with self.subTest(field=field):
                form = ManpowerForecastForm(data=self.data(**{field: value}))
                self.assertIn(field, form.errors)
                self.assertEqual(len(form.errors[field]), 1)

    def test_combined_workload_too_high_blocked(self):
        form = ManpowerForecastForm(data=self.data(growth_rate_percentage="60", aht_change_percentage="60"))
        self.assertIn("together multiply workload", form.errors["aht_change_percentage"][0])

    def test_process_must_belong_to_function(self):
        other = FunctionMaster.objects.create(function_name="HR")
        self.assertIn("process", ManpowerForecastForm(data=self.data(function=other.pk)).errors)


class AutoAhtServiceTests(TestCase):
    def setUp(self):
        self.user = make_user("analyst", "Analyst")
        self.function, self.process = make_masters()

    def test_aht_derived_from_total_time(self):
        f = make_forecast(self.user, self.function, self.process, total_time_spent_hours=Decimal("300"))
        self.assertEqual(f.avg_processing_time, Decimal("8.1818"))
        self.assertEqual(f.time_unit, "MINUTES")
        self.assertEqual(f.required_fte, Decimal("2.01"))      # 100/day x 8.1818 / 408
        self.assertEqual(f.required_people, 3)
        self.assertEqual(f.people_gap, 0)

    def test_aht_recalculated_on_edit(self):
        f = make_forecast(self.user, self.function, self.process, total_time_spent_hours=Decimal("300"))
        f.total_time_spent_hours = Decimal("440")
        save_forecast(f, self.user)
        self.assertEqual(f.avg_processing_time, Decimal("12.0000"))

    def test_missing_total_time_raises(self):
        f = ManpowerForecast(function=self.function, process=self.process, frequency="MONTHLY",
                             volume=Decimal("100"), current_fte=Decimal("1"))
        with self.assertRaises(ValueError):
            save_forecast(f, self.user)

    def test_calculation_steps(self):
        f = make_forecast(self.user, self.function, self.process, total_time_spent_hours=Decimal("300"),
                          frequency="WEEKLY", volume=Decimal("550"), working_days_per_week=Decimal("5.5"))
        steps = {s["label"]: s for s in calculation_steps(f)}
        self.assertEqual(steps["Total time spent"]["result"], "18,000.00 min")
        self.assertEqual(steps["AHT (calculated)"]["formula"], "18,000.00 min / 550.00 items")
        self.assertEqual(steps["AHT (calculated)"]["result"], "32.73 min per item")
        self.assertEqual(steps["Working week"]["formula"], "5.5 days x 8.00 hours")
        self.assertEqual(steps["Working week"]["result"], "44.00 hours per person")
        self.assertEqual(steps["Daily volume"]["formula"], "550.00 / 5.5 days")
        self.assertEqual(steps["Required FTE (whole people)"]["result"], "9 FTE")

    def test_scenarios(self):
        f = make_forecast(self.user, self.function, self.process, run_monte_carlo=False,
                          growth_rate_percentage=Decimal("5"))
        rows = compare_scenarios(build_input(f), [{"name": "Automation", "aht_change_percentage": -2},
                                                  {"name": "Proposed", "current_fte": 4}])
        self.assertEqual(rows[0]["required_people"], 3)
        self.assertEqual(rows[1]["horizon_aht"], 7.85)
        self.assertEqual(rows[2]["fte_gap"], Decimal("1.00"))


def build_upload(rows, headers=None):
    wb = Workbook()
    ws = wb.active
    ws.append(headers or [h for h, _ in excel_service.UPLOAD_COLUMNS])
    for r in rows:
        ws.append(r)
    bio = io.BytesIO()
    wb.save(bio)
    return SimpleUploadedFile("upload.xlsx", bio.getvalue())


def row(function="Finance", process="Invoice Processing", frequency="Monthly", volume=2200, fte=3, hours=300,
        wk=5, mo=22, run_mc="Yes", growth=0, period="Monthly", aht=0, aht_period="Monthly", horizon=12, remarks=""):
    return [function, process, frequency, volume, fte, hours, 8, wk, mo, 15, run_mc, 10, 15, 2000,
            growth, period, aht, aht_period, horizon, remarks]


class ExcelUploadTests(TestCase):
    def setUp(self):
        self.user = make_user("analyst", "Analyst")
        self.function, self.process = make_masters()
        self.vendor = ProcessMaster.objects.create(function=self.function, process_name="Vendor Payments")

    def test_valid_and_invalid_rows(self):
        upload = build_upload([
            row(wk=5.5, mo=23.83),
            row(process="Vendor Payments", fte=2.5),
            row(process="Vendor Payments", hours=None),
            row(function="Unknown"),
        ])
        result = excel_service.process_upload(upload, self.user)
        self.assertEqual(result.success_count, 1)
        errors = {e["row_number"]: e["errors"] for e in result.errors}
        self.assertIn("whole number", errors[3])
        self.assertIn("Total Time Spent Hours: Total time spent (hours) is required", errors[4])
        self.assertIn("does not exist", errors[5])
        f = ManpowerForecast.objects.get()
        self.assertEqual(f.avg_processing_time, Decimal("8.1818"))
        self.assertEqual(f.working_days_per_week, Decimal("5.50"))

    def test_legacy_template_with_aht_column(self):
        headers = ["Function Name", "Process Name", "Frequency", "Volume", "Current FTE",
                   "Average Processing Time", "Time Unit"]
        upload = build_upload([["Finance", "Invoice Processing", "Monthly", 2200, 3, 10, "Minutes"]], headers=headers)
        result = excel_service.process_upload(upload, self.user)
        self.assertEqual(result.success_count, 1, result.errors)
        f = ManpowerForecast.objects.get()
        self.assertEqual(f.total_time_spent_hours, Decimal("366.6667"))
        self.assertEqual(f.required_fte, Decimal("2.45"))

    def test_missing_time_column(self):
        headers = ["Function Name", "Process Name", "Frequency", "Volume", "Current FTE"]
        with self.assertRaises(excel_service.UploadFormatError):
            excel_service.process_upload(build_upload([["Finance", "X", "Daily", 1, 1]], headers=headers), self.user)

    def test_template_and_reports(self):
        wb = load_workbook(io.BytesIO(excel_service.generate_upload_template()))
        headers = [c.value for c in wb["Forecast Upload"][1]]
        self.assertIn("Total Time Spent Hours", headers)
        self.assertNotIn("Average Processing Time", headers)
        make_forecast(self.user, self.function, self.process, growth_rate_percentage=Decimal("5"))
        report = load_workbook(io.BytesIO(excel_service.export_forecasts(ManpowerForecast.objects.all())))
        ws = report["Forecast Report"]
        labels = [c.value for c in ws[2]]
        self.assertEqual(ws.cell(row=3, column=labels.index("AHT (min, calculated)") + 1).value, 10.0)
        self.assertEqual(ws.cell(row=3, column=labels.index("Required FTE (people)") + 1).value, 3)
        self.assertEqual(ws.cell(row=3, column=labels.index("Current FTE") + 1).value, 3)
        detail = load_workbook(io.BytesIO(excel_service.export_forecast_detail(ManpowerForecast.objects.get())))
        self.assertIn("Month-by-Month Projection", detail.sheetnames)
