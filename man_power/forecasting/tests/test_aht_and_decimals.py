"""AHT field + Calculate AHT, decimal Current FTE / Volume, upload and reports."""
import io
from decimal import Decimal

from django.core.files.uploadedfile import SimpleUploadedFile
from django.template import Context, Template
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from forecasting.forms import ManpowerForecastForm, ScenarioForm
from forecasting.models import ManpowerForecast, ProcessMaster
from forecasting.services import excel_service
from forecasting.services.explain import calculation_steps
from forecasting.services.forecast_service import save_forecast

from .helpers import PASSWORD, form_data, make_forecast, make_masters, make_user


class TemplateFilterTests(TestCase):
    def test_fte_dec_filter_registered(self):
        out = Template("{% load forecast_tags %}{{ a|fte_dec }}|{{ b|fte_dec }}|{{ c|fte }}|{{ d|signed_fte }}").render(
            Context({"a": Decimal("2.50"), "b": Decimal("3.00"), "c": Decimal("2.45"), "d": Decimal("-0.50")}))
        self.assertEqual(out, "2.5|3|3|-0.5")


class FormTests(TestCase):
    def setUp(self):
        self.function, self.process = make_masters()

    def test_fields(self):
        fields = ManpowerForecastForm().fields
        self.assertIn("avg_processing_time", fields)
        self.assertNotIn("total_time_spent_hours", fields)

    def test_decimal_fte_and_volume(self):
        form = ManpowerForecastForm(data=form_data(self.function, self.process, current_fte="2.5", volume="6305.5"))
        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual(form.cleaned_data["current_fte"], Decimal("2.5"))
        self.assertEqual(form.cleaned_data["volume"], Decimal("6305.5"))

    def test_negative_fte_rejected(self):
        self.assertIn("current_fte", ManpowerForecastForm(data=form_data(self.function, self.process,
                                                                         current_fte="-0.5")).errors)

    def test_blank_aht_is_calculated_without_contingency(self):
        a = ManpowerForecastForm(data=form_data(self.function, self.process, avg_processing_time="",
                                                contingency_percentage="0"))
        b = ManpowerForecastForm(data=form_data(self.function, self.process, avg_processing_time="",
                                                contingency_percentage="30"))
        self.assertTrue(a.is_valid() and b.is_valid())
        self.assertEqual(a.cleaned_data["avg_processing_time"], Decimal("14.4000"))   # 3 x 8 x 22 x 60 / 2200
        self.assertEqual(a.cleaned_data["avg_processing_time"], b.cleaned_data["avg_processing_time"])

    def test_user_example(self):
        form = ManpowerForecastForm(data=form_data(self.function, self.process, frequency="YEARLY", volume="6305",
                                                   current_fte="180", avg_processing_time=""))
        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual(form.cleaned_data["avg_processing_time"], Decimal("3617.7002"))

    def test_zero_fte_needs_measured_aht(self):
        form = ManpowerForecastForm(data=form_data(self.function, self.process, avg_processing_time="", current_fte="0"))
        self.assertIn("cannot be calculated", form.errors["avg_processing_time"][0])
        self.assertTrue(ManpowerForecastForm(data=form_data(self.function, self.process, current_fte="0")).is_valid())

    def test_scenario_decimal_fte(self):
        self.assertTrue(ScenarioForm(data={"current_fte": "3.5"}).is_valid())


class ServiceTests(TestCase):
    def setUp(self):
        self.user = make_user("analyst", "Analyst")
        self.function, self.process = make_masters()

    def test_entered_aht(self):
        f = make_forecast(self.user, self.function, self.process, avg_processing_time=Decimal("8"), run_monte_carlo=False)
        self.assertEqual(f.avg_processing_time, Decimal("8.0000"))
        self.assertFalse(f.aht_is_calculated)

    def test_calculated_aht(self):
        f = make_forecast(self.user, self.function, self.process, avg_processing_time=None, run_monte_carlo=False)
        self.assertEqual(f.avg_processing_time, Decimal("14.4000"))
        self.assertEqual(f.utilization_percentage, Decimal("117.65"))
        self.assertTrue(f.aht_is_calculated)
        steps = {s["label"]: s for s in calculation_steps(f)}
        self.assertEqual(steps["Team time available in the period"]["formula"],
                         "3 FTE x 8.00 h x 22 days (contingency not deducted)")
        self.assertEqual(steps["AHT (calculated with Calculate AHT)"]["result"], "14.40 min per item")

    def test_decimal_fte_results(self):
        f = make_forecast(self.user, self.function, self.process, current_fte=Decimal("2.5"), run_monte_carlo=False)
        self.assertEqual(f.required_people, 3)
        self.assertEqual(f.people_gap, Decimal("-0.50"))
        self.assertEqual(f.utilization_percentage, Decimal("98.04"))
        self.assertIn("2.5 FTE", f.status_explanation)
        steps = {s["label"]: s for s in calculation_steps(f)}
        self.assertEqual(steps["Current FTE"]["result"], "2.5 FTE")
        self.assertEqual(steps["Gap"]["result"], "-0.5 FTE")

    def test_zero_fte_without_aht_raises(self):
        f = ManpowerForecast(function=self.function, process=self.process, frequency="MONTHLY",
                             volume=Decimal("100"), current_fte=Decimal("0"))
        with self.assertRaises(ValueError):
            save_forecast(f, self.user)


class WebTests(TestCase):
    def setUp(self):
        self.analyst = make_user("analyst", "Analyst")
        self.function, self.process = make_masters()
        self.client.login(username="analyst", password=PASSWORD)

    def test_form_has_button(self):
        r = self.client.get(reverse("forecasting:forecast_create"))
        self.assertContains(r, "data-aht-calc")
        self.assertContains(r, 'step="any"')

    def test_create_decimal_and_result_page(self):
        r = self.client.post(reverse("forecasting:forecast_create"),
                             form_data(self.function, self.process, current_fte="2.5", volume="2200.5",
                                       run_monte_carlo=None), follow=True)
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "AHT (entered)")
        self.assertContains(r, "2.5 FTE")
        f = ManpowerForecast.objects.get()
        self.assertEqual(f.current_fte, Decimal("2.50"))
        edit = self.client.get(reverse("forecasting:forecast_update", args=[f.pk]))
        self.assertContains(edit, 'name="current_fte" value="2.5"')
        self.assertContains(edit, 'name="volume" value="2200.5"')

    def test_create_with_calculated_aht(self):
        r = self.client.post(reverse("forecasting:forecast_create"),
                             form_data(self.function, self.process, avg_processing_time="", run_monte_carlo=None),
                             follow=True)
        self.assertContains(r, "Status not reliable yet")
        f = ManpowerForecast.objects.get()
        self.assertContains(self.client.get(reverse("forecasting:forecast_update", args=[f.pk])), 'data-calculated="true"')


def build_upload(rows, headers=None):
    wb = Workbook()
    ws = wb.active
    ws.append(headers or [h for h, _ in excel_service.UPLOAD_COLUMNS])
    for r in rows:
        ws.append(r)
    bio = io.BytesIO()
    wb.save(bio)
    return SimpleUploadedFile("upload.xlsx", bio.getvalue())


def row(process="Invoice Processing", fte=3, aht=10, volume=2200):
    return ["Finance", process, "Monthly", volume, fte, aht, 8, 5, 22, 15, None, "No", None, None, None,
            0, "Monthly", 0, "Monthly", 12, ""]


class UploadAndReportTests(TestCase):
    def setUp(self):
        self.user = make_user("analyst", "Analyst")
        self.function, self.process = make_masters()
        ProcessMaster.objects.create(function=self.function, process_name="Vendor Payments")
        ProcessMaster.objects.create(function=self.function, process_name="GST")

    def test_upload(self):
        result = excel_service.process_upload(build_upload([
            row(fte=2.5, volume=2200.5), row(process="Vendor Payments", aht=None), row(process="GST", aht=None, fte=0),
        ]), self.user)
        self.assertEqual(result.success_count, 2, result.errors)
        self.assertEqual(result.auto_filled_rows, [3])
        self.assertIn("cannot be filled", result.errors[0]["errors"])
        f = ManpowerForecast.objects.get(process=self.process)
        self.assertEqual(f.current_fte, Decimal("2.50"))
        self.assertEqual(f.avg_processing_time, Decimal("10.0000"))
        self.assertEqual(ManpowerForecast.objects.get(process__process_name="Vendor Payments").avg_processing_time,
                         Decimal("14.4000"))

    def test_legacy_aht_column(self):
        headers = ["Function Name", "Process Name", "Frequency", "Volume", "Current FTE", "Average Processing Time", "Time Unit"]
        result = excel_service.process_upload(build_upload([["Finance", "Invoice Processing", "Monthly", 2200, 3, 10, "Minutes"]],
                                                           headers=headers), self.user)
        self.assertEqual(result.success_count, 1, result.errors)

    def test_reports(self):
        make_forecast(self.user, self.function, self.process, avg_processing_time=None, current_fte=Decimal("2.5"))
        ws = load_workbook(io.BytesIO(excel_service.export_forecasts(ManpowerForecast.objects.all())))["Forecast Report"]
        labels = [c.value for c in ws[2]]
        self.assertEqual(ws.cell(row=3, column=labels.index("Current FTE") + 1).value, 2.5)
        self.assertEqual(ws.cell(row=3, column=labels.index("AHT Source") + 1).value, "Calculated (Calculate AHT)")
        tpl = load_workbook(io.BytesIO(excel_service.generate_upload_template()))
        self.assertIn("AHT Minutes", [c.value for c in tpl["Forecast Upload"][1]])
